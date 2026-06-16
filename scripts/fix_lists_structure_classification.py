from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph


DOCX_PATH = Path(sys.argv[1])
ROOT = Path(sys.argv[2])
OUTPUT_PATH = Path(sys.argv[3]) if len(sys.argv) > 3 else DOCX_PATH


HEADING_MAP = {
    "1. Анализ предметной области": "1. Аналитическая предметной области",
    "1.1. Сведения об организации": "1.1 Сведения об организации",
    "1.2. Информационная и функциональная структура": "1.2 Информационная и функциональная структура",
    "1.3. Классификация информационных систем": "1.3 Классификация веб-сайтов",
    "1.3. Классификация веб-сайтов": "1.3 Классификация веб-сайтов",
    "1.4. Обоснование выбора программных средств для создания системы заявок": "1.4 Обоснование выбора программных средств для создания системы заявок",
    "1.5. Постановка задачи": "1.5 Постановка задачи",
    "2.1. Проектирование архитектуры системы заявок": "2.1 Проектирование архитектуры системы заявок",
    "2.2. Разработка интерфейсов": "2.2 Разработка интерфейсов",
    "2.3. Разработка клиентской части информационной системы": "2.3 Разработка клиентской части системы заявок",
    "2.4. Разработка серверной части информационной системы": "2.4 Разработка серверной части системы заявок",
    "2.5. Тестирование информационной системы": "2.6 Тестирование системы заявок",
    "3. Экономическое обоснование": "3. Обоснование экономической эффективности проекта",
    "ПРИЛОЖЕНИЕ 1. Код для функционирования системы заявок": "Приложение 1  Код для функционирования системы заявок",
}

CLASSIFICATION_PARAS = [
    "Веб-сайты классифицируются по назначению, составу функций и способу взаимодействия с пользователем. Для дипломного проекта это важно, потому что разрабатываемая система заявок не является обычной информационной страницей: она относится к интерактивным web-приложениям, где пользователь вводит данные, проходит авторизацию и получает результат обработки запроса.",
    "Основные виды веб-сайтов:",
    "- сайт-визитка, содержащий краткую информацию об организации, контакты и основные сведения о деятельности;",
    "- корпоративный сайт, предназначенный для представления компании, публикации новостей, структуры, услуг и внутренних материалов;",
    "- информационный портал, объединяющий большое количество материалов, разделов, новостей и справочной информации;",
    "- интернет-магазин, обеспечивающий каталог товаров, корзину, оформление заказа и взаимодействие с покупателем;",
    "- лендинг, ориентированный на представление одного продукта, услуги или предложения и получение целевого действия от пользователя;",
    "- блог или новостной сайт, где основное содержание формируется в виде регулярно публикуемых материалов;",
    "- web-приложение или web-сервис, в котором сайт выполняет прикладную функцию: авторизацию, обработку данных, хранение записей, фильтрацию, управление ролями и формирование результата.",
    "Разрабатываемая система заявок относится к типу web-приложения. Она не ограничивается публикацией статической информации, а предоставляет пользователям рабочий инструмент для регистрации и обработки обращений. В отличие от сайта-визитки или лендинга, такая система имеет базу данных, ролевую модель, формы ввода, серверную проверку прав и динамическое отображение записей.",
]


def set_text(p: Paragraph, text: str) -> None:
    for run in p.runs:
        run.text = ""
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)


def find_para(doc: Document, needle: str) -> Paragraph:
    for p in doc.paragraphs:
        if needle in p.text:
            return p
    raise ValueError(needle)


def paragraph_after(paragraph: Paragraph, text: str = "", style=None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style is not None:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def delete_paragraph(p: Paragraph) -> None:
    element = p._element
    element.getparent().remove(element)
    p._p = p._element = None


def clear_between(doc: Document, start: Paragraph, end: Paragraph) -> None:
    body = doc._body._element
    children = list(body)
    start_i = children.index(start._p)
    end_i = children.index(end._p)
    for element in children[start_i + 1:end_i]:
        body.remove(element)


def add_picture_after(paragraph: Paragraph, image: Path, caption: str, width_cm: float = 13.0) -> Paragraph:
    p_img = paragraph_after(paragraph)
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.add_run().add_picture(str(image), width=Cm(width_cm))

    p_cap = paragraph_after(p_img, caption)
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p_cap.runs:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
    return p_cap


def replace_classification(doc: Document) -> None:
    start = find_para(doc, "1.3")
    end = find_para(doc, "1.4")
    clear_between(doc, start, end)
    current = start
    for text in CLASSIFICATION_PARAS:
        p = paragraph_after(current, text)
        p.paragraph_format.first_line_indent = Cm(1.25) if not text.startswith("- ") else Cm(0)
        current = p
    current = add_picture_after(
        current,
        ROOT / "docs" / "mockup_images" / "mockup_login.png",
        "Рисунок 1.4 – Пример web-приложения: экран авторизации пользователя",
    )
    current = add_picture_after(
        current,
        ROOT / "docs" / "mockup_images" / "mockup_journal.png",
        "Рисунок 1.5 – Пример web-приложения: рабочий журнал заявок",
    )
    p = paragraph_after(current, "На рисунках 1.4 и 1.5 показаны экраны web-приложения, выбранного в качестве типа разрабатываемого сайта: пользователь взаимодействует с формами, а содержимое страниц формируется динамически на основе данных системы.")
    p.paragraph_format.first_line_indent = Cm(1.25)


def normalize_headings(doc: Document) -> None:
    for p in doc.paragraphs:
        text = p.text.strip()
        if text in HEADING_MAP:
            set_text(p, HEADING_MAP[text])


def is_heading_or_caption(p: Paragraph) -> bool:
    style = getattr(p.style, "name", "")
    text = p.text.strip()
    if style.startswith("Heading"):
        return True
    if re.match(r"^(Рисунок|Таблица|Листинг)\s+\d+", text):
        return True
    if text.startswith("Приложение "):
        return True
    return False


def is_bibliography_item(text: str, in_bibliography: bool) -> bool:
    return in_bibliography and re.match(r"^\d+\.\s+", text) is not None


def convert_numbered_lists(doc: Document) -> None:
    in_bibliography = False
    i = 0
    paragraphs = doc.paragraphs
    while i < len(paragraphs):
        text = paragraphs[i].text.strip()
        if text == "Библиография":
            in_bibliography = True
        elif text.startswith("Приложение 1"):
            in_bibliography = False

        if in_bibliography or is_heading_or_caption(paragraphs[i]) or not re.match(r"^\d+\.\s+", text):
            i += 1
            continue

        group: list[Paragraph] = []
        while i < len(paragraphs):
            current = paragraphs[i]
            current_text = current.text.strip()
            if is_heading_or_caption(current) or in_bibliography or not re.match(r"^\d+\.\s+", current_text):
                break
            group.append(current)
            i += 1

        for index, p in enumerate(group):
            item = re.sub(r"^\d+\.\s*", "", p.text.strip()).strip()
            item = item.rstrip(" ;.")
            ending = "." if index == len(group) - 1 else ";"
            set_text(p, f"- {item}{ending}")
    # Static Word list numbering can survive text replacement; remove list properties.
    for p in doc.paragraphs:
        if p.text.strip().startswith("- "):
            p.style = doc.styles["Normal"]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(0.75)


def fix_existing_dash_lists(doc: Document) -> None:
    paragraphs = doc.paragraphs
    i = 0
    while i < len(paragraphs):
        if not paragraphs[i].text.strip().startswith("- "):
            i += 1
            continue
        group = []
        while i < len(paragraphs) and paragraphs[i].text.strip().startswith("- "):
            group.append(paragraphs[i])
            i += 1
        for index, p in enumerate(group):
            text = p.text.strip()[2:].strip().rstrip(" ;.")
            ending = "." if index == len(group) - 1 else ";"
            set_text(p, f"- {text}{ending}")


def force_black(doc: Document) -> None:
    for p in doc.paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor(0, 0, 0)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.color.rgb = RGBColor(0, 0, 0)


def main() -> None:
    doc = Document(str(DOCX_PATH))
    normalize_headings(doc)
    replace_classification(doc)
    normalize_headings(doc)
    convert_numbered_lists(doc)
    fix_existing_dash_lists(doc)
    force_black(doc)
    doc.save(str(OUTPUT_PATH))
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
