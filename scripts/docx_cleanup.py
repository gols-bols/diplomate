from __future__ import annotations

import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


def clear_paragraph_breaks(paragraph: Paragraph) -> bool:
    changed = False
    p_pr = paragraph._p.find(qn("w:pPr"))
    if p_pr is not None:
        for tag in ("w:pageBreakBefore", "w:columnBreakBefore"):
            node = p_pr.find(qn(tag))
            if node is not None:
                p_pr.remove(node)
                changed = True
        sect = p_pr.find(qn("w:sectPr"))
        if sect is not None:
            p_pr.remove(sect)
            changed = True
        for tag in ("w:keepNext", "w:keepLines", "w:pageBreakAfter"):
            node = p_pr.find(qn(tag))
            if node is not None:
                p_pr.remove(node)
                changed = True
    for br in paragraph._p.xpath(".//w:br"):
        if br.get(qn("w:type")) in ("page", "column"):
            br.getparent().remove(br)
            changed = True
    for node in paragraph._p.xpath(".//w:lastRenderedPageBreak"):
        node.getparent().remove(node)
        changed = True
    paragraph.paragraph_format.page_break_before = False
    return changed


def sanitize_paragraph_element(element) -> None:
    p_pr = element.find(qn("w:pPr"))
    if p_pr is not None:
        for tag in ("w:pageBreakBefore", "w:sectPr", "w:keepNext", "w:keepLines"):
            node = p_pr.find(qn(tag))
            if node is not None:
                p_pr.remove(node)
    for br in element.xpath(".//w:br"):
        if br.get(qn("w:type")) in ("page", "column"):
            br.getparent().remove(br)
    for node in element.xpath(".//w:lastRenderedPageBreak"):
        node.getparent().remove(node)


def delete_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def find_section_range(doc: Document, start_pattern: str, end_pattern: str) -> tuple[int, int] | None:
    start = end = None
    for i, paragraph in enumerate(doc.paragraphs):
        text = (paragraph.text or "").strip()
        if start is None and re.match(start_pattern, text):
            start = i
            continue
        if start is not None and re.match(end_pattern, text):
            end = i
            break
    if start is None:
        return None
    return start, end or len(doc.paragraphs)


def insert_paragraph_after(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def add_body_paragraph_after(doc: Document, anchor: Paragraph, text: str) -> Paragraph:
    new_p = insert_paragraph_after(anchor)
    try:
        new_p.style = doc.styles["Normal"]
    except KeyError:
        new_p.style = doc.styles["Обычный"] if "Обычный" in [s.name for s in doc.styles] else doc.styles["Normal"]
    fmt = new_p.paragraph_format
    fmt.line_spacing = 1.5
    fmt.first_line_indent = Cm(1.25)
    fmt.left_indent = Cm(0)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.page_break_before = False
    run = new_p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    clear_paragraph_breaks(new_p)
    return new_p


def fix_chapter_heading_breaks(doc: Document) -> int:
    """Keep page breaks only before major Heading 1 (chapters, conclusion, appendices)."""
    allowed = 0
    for paragraph in doc.paragraphs:
        text = (paragraph.text or "").strip()
        style = paragraph.style.name if paragraph.style else ""
        if style != "Heading 1":
            clear_paragraph_breaks(paragraph)
            continue
        keep = bool(
            re.match(r"^[123]\.\s", text)
            or text in {"Введение", "Заключение", "ЗАКЛЮЧЕНИЕ", "Библиография", "БИБЛИОГРАФИЯ", "Содержание"}
            or text.startswith(("ПРИЛОЖЕНИЕ", "Приложение"))
        )
        if keep:
            allowed += 1
        else:
            clear_paragraph_breaks(paragraph)
    return allowed


def rebuild_requirements_list(doc: Document) -> bool:
    """Rebuild 1.4 list paragraphs without hidden breaks (fixes pagination after item 13)."""
    span = find_section_range(doc, r"^1\.4[\.\s]", r"^1\.5[\.\s]")
    if span is None:
        return False
    start, end = span
    paragraphs = doc.paragraphs
    header = paragraphs[start]

    intro_lines: list[str] = []
    numbered: list[str] = []
    task_intro: str | None = None
    tasks: list[str] = []
    tail: list[str] = []
    phase = "intro"

    for i in range(start + 1, end):
        text = (paragraphs[i].text or "").strip()
        if not text:
            continue
        if phase == "intro":
            if re.match(r"^1\.\s", text):
                phase = "reqs"
                numbered.append(text)
            else:
                intro_lines.append(text)
            continue
        if phase == "reqs":
            m_num = re.match(r"^(\d+)\.\s", text)
            if m_num and numbered and int(m_num.group(1)) == 1 and any(
                x.lstrip().startswith("13.") for x in numbered
            ):
                phase = "tasks"
                tasks.append(text)
            elif m_num:
                numbered.append(text)
            elif "реализации" in text.lower():
                task_intro = text
                phase = "tasks"
            else:
                if numbered:
                    numbered[-1] = f"{numbered[-1]} {text}"
                elif intro_lines:
                    intro_lines[-1] = f"{intro_lines[-1]} {text}"
            continue
        if phase == "tasks":
            if re.match(r"^\d+\.\s", text):
                tasks.append(text)
            elif task_intro is None:
                task_intro = text
            elif tasks and not re.match(r"^\d+\.\s", text):
                tail.append(text)
            else:
                if tasks:
                    tasks[-1] = f"{tasks[-1]} {text}"
                else:
                    task_intro = f"{task_intro} {text}"
            continue

    for i in range(end - 1, start, -1):
        delete_paragraph(paragraphs[i])

    anchor = header
    for line in intro_lines:
        anchor = add_body_paragraph_after(doc, anchor, line)
    for line in numbered:
        anchor = add_body_paragraph_after(doc, anchor, line)
    if task_intro:
        anchor = add_body_paragraph_after(doc, anchor, task_intro)
    for line in tasks:
        anchor = add_body_paragraph_after(doc, anchor, line)
    for line in tail:
        anchor = add_body_paragraph_after(doc, anchor, line)

    return True
