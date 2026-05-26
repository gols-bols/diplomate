from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from docx_cleanup import fix_chapter_heading_breaks, rebuild_requirements_list
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx_figures import FigureCounter, insert_figures_after_anchor, max_figure_number
from project_paths import MOCKUP_FIGURES
from resolve_diploma_v2 import resolve_diploma_v2


def _insert_paragraph_after(paragraph, text: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    p = Paragraph(new_p, paragraph._parent)
    if text:
        p.add_run(text)
    return p


def _set_body_paragraph(p) -> None:
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = p.paragraph_format
    fmt.line_spacing = 1.5
    fmt.first_line_indent = Cm(1.25)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)


def _set_code_paragraph(p) -> None:
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = p.paragraph_format
    fmt.line_spacing = 1.0
    fmt.first_line_indent = Cm(0)
    fmt.left_indent = Cm(0)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)


def _add_code_block(doc: Document, title: str, code: str) -> None:
    h = doc.add_paragraph(title)
    h.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    hr = h.runs[0] if h.runs else h.add_run(title)
    hr.font.name = "Times New Roman"
    hr.font.size = Pt(14)
    hr.bold = True

    for line in code.rstrip().splitlines():
        p = doc.add_paragraph()
        _set_code_paragraph(p)
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(10)


def _insert_code_block_after(anchor: Paragraph, title: str, code: str) -> Paragraph:
    """Insert code block paragraphs right after anchor; returns last paragraph."""
    h = _insert_paragraph_after(anchor, title)
    h.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if h.runs:
        hr = h.runs[0]
        hr.font.name = "Times New Roman"
        hr.font.size = Pt(14)
        hr.bold = True
    last = h
    for line in code.rstrip().splitlines():
        p = _insert_paragraph_after(last, "")
        _set_code_paragraph(p)
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(10)
        last = p
    return last


def _find_heading(doc: Document, pattern: re.Pattern[str]) -> int | None:
    for i, p in enumerate(doc.paragraphs):
        t = (p.text or "").strip()
        if pattern.match(t):
            return i
    return None


def _delete_range(doc: Document, start: int, end_inclusive: int) -> None:
    for i in range(end_inclusive, start - 1, -1):
        el = doc.paragraphs[i]._element
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)


def _ensure_body_after_each_heading(doc: Document) -> int:
    """No naked numbered headings in chapters 1–2."""
    added = 0
    for i in range(len(doc.paragraphs) - 1):
        p = doc.paragraphs[i]
        t = (p.text or "").strip()
        style = p.style.name if p.style else ""
        if not style.startswith("Heading"):
            continue
        if not re.match(r"^[12]\.\d", t):
            continue

        # find next non-empty paragraph
        j = i + 1
        while j < len(doc.paragraphs) and not (doc.paragraphs[j].text or "").strip():
            j += 1
        if j >= len(doc.paragraphs):
            body = p.insert_paragraph_after("Текст по данному пункту приведён в последующих подпунктах.")
            _set_body_paragraph(body)
            added += 1
            continue

        next_style = doc.paragraphs[j].style.name if doc.paragraphs[j].style else ""
        if next_style.startswith("Heading"):
            body = p.insert_paragraph_after("Описание и результаты по данному подпункту приведены далее по тексту.")
            _set_body_paragraph(body)
            added += 1

    return added


def _restore_comparison_section(doc: Document) -> bool:
    """Insert 'обзор аналогов' + 'сравнение' text under 1.2 if missing."""
    needles = ("GLPI", "Jira", "ServiceDesk", "сравнение аналогов")
    text_all = "\n".join((p.text or "") for p in doc.paragraphs).lower()
    if any(n.lower() in text_all for n in needles):
        return False

    idx = _find_heading(doc, re.compile(r"^1\.2\.\s"))
    if idx is None:
        return False

    insert_after = doc.paragraphs[idx]
    blocks = [
        "Анализ аналогов необходим для обоснования собственной реализации. В качестве типовых аналогов рассматриваются: "
        "табличный учет обращений, коммуникационные решения (почта/мессенджеры) и профессиональные helpdesk-системы.",
        "Сравнение аналогов показывает, что для целей дипломной работы наиболее рациональным является создание собственного "
        "web-приложения с прикладно достаточным набором функций. В отличие от табличных и коммуникационных решений, оно "
        "обеспечивает централизованное хранение данных. В отличие от тяжелых платформ, оно остается прозрачным, управляемым "
        "и пригодным для поэтапного развития.",
    ]
    for block in reversed(blocks):
        p = insert_after.insert_paragraph_after(block)
        _set_body_paragraph(p)
    return True


def _rewrite_appendices(doc: Document, repo_root: Path) -> None:
    """Appendix 1 = code listing, Appendix 2 = interface screenshots."""
    app1 = _find_heading(doc, re.compile(r"^ПРИЛОЖЕНИЕ\s+1\b", re.I))
    app2 = _find_heading(doc, re.compile(r"^ПРИЛОЖЕНИЕ\s+2\b", re.I))
    if app1 is None or app2 is None or app2 <= app1:
        return

    # Update headings
    doc.paragraphs[app1].text = "ПРИЛОЖЕНИЕ 1. Код для функционирования системы заявок"
    doc.paragraphs[app2].text = "ПРИЛОЖЕНИЕ 2. Интерфейс системы заявок (скриншоты)"

    # Clear content inside appendices
    _delete_range(doc, app1 + 1, app2 - 1)
    _delete_range(doc, app2 + 1, len(doc.paragraphs) - 1)

    # Insert code listing into Appendix 1
    intro = _insert_paragraph_after(
        doc.paragraphs[app1],
        "В приложении приведены ключевые фрагменты исходного кода (маршруты, контроллер и модель), "
        "обеспечивающие работу системы управления заявками.",
    )
    _set_body_paragraph(intro)

    def read_file(rel: str, max_lines: int = 120) -> str:
        p = repo_root / rel
        if not p.is_file():
            return f"// Файл не найден: {rel}"
        lines = p.read_text(encoding="utf-8", errors="replace").splitlines()
        if len(lines) > max_lines:
            lines = lines[:max_lines] + ["// ... (обрезано) ..."]
        return "\n".join(lines)

    # Insert code blocks inside Appendix 1 (between Appendix 1 and Appendix 2 headings).
    cur = intro
    cur = _insert_code_block_after(cur, "Листинг 1 — routes/web.php", read_file("routes/web.php", 80))
    cur = _insert_code_block_after(cur, "Листинг 2 — routes/api.php", read_file("routes/api.php", 60))
    cur = _insert_code_block_after(
        cur,
        "Листинг 3 — app/Http/Controllers/TicketController.php",
        read_file("app/Http/Controllers/TicketController.php", 140),
    )
    _insert_code_block_after(cur, "Листинг 4 — app/Models/Ticket.php", read_file("app/Models/Ticket.php", 120))

    # Insert interface figures into Appendix 2 (after its heading)
    # Find updated app2 again (indices changed after deletions/additions); search by exact prefix.
    app2_new = _find_heading(doc, re.compile(r"^ПРИЛОЖЕНИЕ\s+2\.", re.I))
    if app2_new is None:
        return
    anchor = doc.paragraphs[app2_new]
    counter = FigureCounter(max_figure_number(doc))
    insert_figures_after_anchor(doc, anchor, list(MOCKUP_FIGURES), counter=counter)


def verify_media(path: Path) -> int:
    from zipfile import ZipFile

    with ZipFile(path) as zf:
        return len([n for n in zf.namelist() if n.startswith("word/media/")])


def main() -> None:
    repo_root = Path(__file__).resolve().parent.parent

    target = resolve_diploma_v2()
    backup = target.with_name(target.stem + ".bak_final.docx")
    final = target.with_name(target.stem + "_ФИНАЛ.docx")
    shutil.copy2(target, backup)

    doc = Document(str(target))

    rebuild_requirements_list(doc)
    fix_chapter_heading_breaks(doc)
    _restore_comparison_section(doc)
    _ensure_body_after_each_heading(doc)
    _rewrite_appendices(doc, repo_root)

    doc.save(str(target))
    doc = Document(str(target))
    doc.save(str(target))  # second save for compatibility

    shutil.copy2(target, final)

    media_n = verify_media(target)
    if media_n < 10:
        raise SystemExit(f"FAIL: expected >=10 images, got {media_n}")

    # Avoid printing paths with unicode-normalization issues in Windows console.
    print("OK")
    print("media:", media_n)


if __name__ == "__main__":
    main()

