from __future__ import annotations

import argparse
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from docx_figures import (
    FigureCounter,
    add_figures_before_anchor,
    max_figure_number,
    remove_mockup_block_before,
    set_paragraph_text,
    set_run_font,
)
from project_paths import DOCS_DIR, MOCKUP_FIGURES, PROJECT_ROOT

# По заданию ДП: макеты интерфейса — в п. 2.2 «Разработка интерфейсов» (перед 2.3).
TARGETS = [
    (
        DOCS_DIR / "Диплом_отдельный_финал_60.docx",
        re.compile(r"^2\.3\s"),
        "2.2.3. Скриншоты макетов интерфейса",
    ),
    (
        DOCS_DIR / "Пояснительная_записка_финал.docx",
        re.compile(r"^2\.3\s"),
        "Скриншоты макетов интерфейса",
    ),
]


def find_anchor(doc: Document, pattern: re.Pattern[str]):
    for paragraph in doc.paragraphs:
        text = (paragraph.text or "").strip()
        if pattern.match(text):
            return paragraph
    return None


def insert_section(doc: Document, anchor, heading: str, counter: FigureCounter) -> int:
    heading_p = anchor.insert_paragraph_before()
    heading_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = heading_p.paragraph_format
    fmt.line_spacing = 1.5
    fmt.first_line_indent = Cm(0)
    fmt.space_after = Pt(0)
    run = heading_p.add_run(heading)
    set_run_font(run, 15, True)

    intro_p = anchor.insert_paragraph_before()
    set_paragraph_text(
        intro_p,
        "В соответствии с заданием на дипломную работу (п. 2.2) ниже приведены скриншоты макетов "
        "основных экранов веб-приложения: авторизация, журнал заявок, создание и карточка обращения.",
        first_indent=True,
    )
    return add_figures_before_anchor(doc, anchor, MOCKUP_FIGURES, counter=counter)


def process_doc(path, anchor_pattern: re.Pattern[str], heading: str, *, force: bool) -> None:
    if not path.is_file():
        print(f"Файл не найден: {path}")
        return

    doc = Document(path)
    anchor = find_anchor(doc, anchor_pattern)
    if anchor is None:
        print(f"Якорь «2.3 …» не найден в {path.name}; рисунки не вставлены")
        return

    removed = remove_mockup_block_before(anchor) if force else 0
    if removed:
        print(f"{path.name}: удалён старый блок ({removed} абзацев)")

    counter = FigureCounter(max_figure_number(doc))
    start_num = counter._current + 1
    count = insert_section(doc, anchor, heading, counter)
    doc.save(str(path.resolve()))
    end_num = counter._current
    print(
        f"{path.name}: вставлено рисунков — {count} "
        f"(нумерация: Рисунок {start_num}–Рисунок {end_num})"
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Вставка макетов интерфейса в финальные docx")
    parser.add_argument(
        "--force",
        action="store_true",
        help="Удалить прежний блок макетов и вставить заново с актуальной нумерацией",
    )
    args = parser.parse_args()

    print(f"Корень проекта: {PROJECT_ROOT}")
    for path, pattern, heading in TARGETS:
        process_doc(path, pattern, heading, force=args.force)


if __name__ == "__main__":
    main()
