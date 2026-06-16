from __future__ import annotations

import sys
from pathlib import Path

from docx import Document


def find(doc: Document, needle: str):
    for p in doc.paragraphs:
        if needle in p.text:
            return p
    raise ValueError(needle)


def delete(p):
    element = p._element
    element.getparent().remove(element)
    p._p = p._element = None


def move_paragraph_before(doc: Document, paragraph_text: str, before_text: str) -> None:
    p = find(doc, paragraph_text)
    before = find(doc, before_text)
    body = doc._body._element
    body.remove(p._p)
    target = list(body).index(before._p)
    body.insert(target, p._p)


def main() -> None:
    path = Path(sys.argv[1])
    doc = Document(str(path))
    move_paragraph_before(
        doc,
        "1.3. Классификация информационных систем",
        "Для сопоставления проектного решения",
    )
    for p in list(doc.paragraphs):
        if p.text.startswith("Для разрабатываемого решения выбран класс внутренней"):
            delete(p)
            break
    doc.save(str(path))
    print(path)


if __name__ == "__main__":
    main()
