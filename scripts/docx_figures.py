from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

from project_paths import PROJECT_ROOT

FIGURE_CAPTION_RE = re.compile(r"Рисунок\s+(\d+)", re.IGNORECASE)


def set_run_font(run, size: int = 14, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def set_paragraph_text(
    paragraph,
    text: str,
    *,
    center: bool = False,
    size: int = 14,
    first_indent: bool = True,
) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    set_run_font(run, size)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.5
    fmt.first_line_indent = Cm(0 if center or not first_indent else 1.25)
    fmt.space_after = Pt(0)


def figure_caption(number: int) -> str:
    return f"Рисунок {number}"


def resolve_image_path(image_path: Path) -> Path | None:
    candidates = [image_path]
    if not image_path.is_absolute():
        candidates.append(PROJECT_ROOT / image_path)
    try:
        candidates.append((PROJECT_ROOT / image_path).resolve())
    except OSError:
        pass
    seen: set[Path] = set()
    for candidate in candidates:
        try:
            resolved = candidate.resolve()
        except OSError:
            resolved = candidate
        if resolved in seen:
            continue
        seen.add(resolved)
        if resolved.is_file():
            return resolved
    return None


def paragraph_has_image(paragraph: Paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing"))


def max_figure_number(doc: Document) -> int:
    maximum = 0
    for paragraph in doc.paragraphs:
        match = FIGURE_CAPTION_RE.search(paragraph.text or "")
        if match:
            maximum = max(maximum, int(match.group(1)))
    return maximum


class FigureCounter:
    def __init__(self, start: int = 0) -> None:
        self._current = start

    def next(self) -> int:
        self._current += 1
        return self._current


def insert_paragraph_before(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    return Paragraph(new_p, paragraph._parent)


def delete_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def is_figure_block_paragraph(paragraph: Paragraph) -> bool:
    text = (paragraph.text or "").strip()
    lower = text.lower()
    if paragraph_has_image(paragraph):
        return True
    if FIGURE_CAPTION_RE.match(text):
        return True
    if lower.startswith("рисунок") and ("—" in text or "–" in text):
        return True
    if lower.startswith("на рисунке"):
        return True
    if text.upper().startswith("ПРИЛОЖЕНИЕ") or text.startswith("Приложение"):
        return False
    if "2.2.3" in lower and "макет" in lower:
        return True
    if lower.startswith("ниже приведены") and ("рисунк" in lower or "скриншот" in lower or "диаграм" in lower):
        return True
    return False


def is_mockup_block_paragraph(paragraph: Paragraph) -> bool:
    text = (paragraph.text or "").strip()
    lower = text.lower()
    if paragraph_has_image(paragraph):
        return True
    if "скриншоты макет" in lower or "2.2.3" in lower:
        return True
    if lower.startswith("ниже приведены скриншоты"):
        return True
    if FIGURE_CAPTION_RE.match(text):
        return True
    if lower.startswith("на рисунке"):
        return True
    if " - макет " in lower and FIGURE_CAPTION_RE.search(text):
        return True
    return False


def remove_figure_block_before(anchor: Paragraph, *, predicate=None) -> int:
    removed = 0
    previous = anchor._p.getprevious()
    while previous is not None:
        paragraph = Paragraph(previous, anchor._parent)
        ok = predicate(paragraph) if predicate else is_figure_block_paragraph(paragraph)
        if not ok:
            break
        previous = previous.getprevious()
        delete_paragraph(paragraph)
        removed += 1
    return removed


def remove_all_figures_from_document(doc: Document) -> int:
    removed = 0
    for paragraph in reversed(doc.paragraphs):
        if is_figure_block_paragraph(paragraph):
            delete_paragraph(paragraph)
            removed += 1
    return removed


def remove_mockup_block_before(anchor: Paragraph) -> int:
    return remove_figure_block_before(anchor, predicate=is_mockup_block_paragraph)


def insert_paragraph_after(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def insert_figure_after(
    doc: Document,
    anchor: Paragraph,
    image_path: Path,
    description: str,
    *,
    counter: FigureCounter,
) -> Paragraph | None:
    resolved = resolve_image_path(image_path)
    if resolved is None:
        print(f"ОШИБКА: файл не найден — {image_path}")
        return None

    number = counter.next()
    caption = figure_caption(number)
    body = description.format(n=number) if "{n}" in description else description

    picture_p = insert_paragraph_after(anchor)
    picture_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_p.add_run().add_picture(str(resolved), width=Cm(16.2))

    cap_p = insert_paragraph_after(picture_p)
    set_paragraph_text(cap_p, caption, center=True, size=13, first_indent=False)

    last = cap_p
    if body.strip():
        desc_p = insert_paragraph_after(cap_p)
        set_paragraph_text(desc_p, body, size=14, first_indent=True)
        last = desc_p
    return last


def insert_figures_after_anchor(
    doc: Document,
    anchor: Paragraph,
    figures: list[tuple[Path, str, str]],
    *,
    counter: FigureCounter | None = None,
) -> int:
    if counter is None:
        counter = FigureCounter(max_figure_number(doc))
    cur = anchor
    inserted = 0
    for image_path, _title, description in figures:
        nxt = insert_figure_after(doc, cur, image_path, description, counter=counter)
        if nxt is not None:
            cur = nxt
            inserted += 1
    return inserted


def add_figure(
    doc: Document,
    image_path: Path,
    description: str,
    *,
    counter: FigureCounter | None = None,
    caption_below: bool = True,
) -> bool:
    resolved = resolve_image_path(image_path)
    if resolved is None:
        print(f"ОШИБКА: файл изображения не найден — {image_path}")
        return False

    if counter is None:
        counter = FigureCounter(max_figure_number(doc))
    number = counter.next()
    caption = figure_caption(number)
    body = description.format(n=number) if "{n}" in description else description

    picture_paragraph = doc.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = picture_paragraph.add_run()
    run.add_picture(str(resolved), width=Cm(16.2))

    caption_paragraph = doc.add_paragraph()
    set_paragraph_text(caption_paragraph, caption, center=True, size=13, first_indent=False)

    if body.strip():
        description_paragraph = doc.add_paragraph()
        set_paragraph_text(description_paragraph, body, size=14, first_indent=True)

    return True


def add_figures_before_anchor(
    doc: Document,
    anchor: Paragraph,
    figures: list[tuple[Path, str, str]],
    *,
    counter: FigureCounter | None = None,
) -> int:
    if counter is None:
        counter = FigureCounter(max_figure_number(doc))

    inserted = 0
    for image_path, _title, description in figures:
        resolved = resolve_image_path(image_path)
        if resolved is None:
            print(f"ОШИБКА: файл изображения не найден — {image_path}")
            continue

        number = counter.next()
        caption = figure_caption(number)
        body = description.format(n=number) if "{n}" in description else description

        block: list[Paragraph] = []
        for _ in range(3):
            block.append(insert_paragraph_before(anchor))
        picture_paragraph, caption_paragraph, description_paragraph = block

        picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = picture_paragraph.add_run()
        run.add_picture(str(resolved), width=Cm(16.2))
        set_paragraph_text(caption_paragraph, caption, center=True, size=13, first_indent=False)
        set_paragraph_text(description_paragraph, body, size=14, first_indent=True)
        inserted += 1

    return inserted
