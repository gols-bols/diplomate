from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph


def find(doc: Document, needle: str) -> Paragraph:
    for p in doc.paragraphs:
        if needle in p.text:
            return p
    raise ValueError(needle)


def paragraph_after(paragraph: Paragraph, text: str = "") -> Paragraph:
    from docx.oxml import OxmlElement

    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    return new_para


def delete(p: Paragraph) -> None:
    element = p._element
    element.getparent().remove(element)
    p._p = p._element = None


def method(source: str, signature: str) -> str:
    start = source.index(signature)
    brace = source.index("{", start)
    depth = 0
    for i in range(brace, len(source)):
        if source[i] == "{":
            depth += 1
        elif source[i] == "}":
            depth -= 1
            if depth == 0:
                return source[start:i + 1]
    raise ValueError(signature)


def add_code_after(paragraph: Paragraph, code: str) -> Paragraph:
    current = paragraph
    for line in code.splitlines():
        p = paragraph_after(current, line)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        for r in p.runs:
            r.font.name = "Courier New"
            r.font.size = Pt(9)
        current = p
    return current


def main() -> None:
    doc_path = Path(sys.argv[1])
    source_path = Path(sys.argv[2])
    doc = Document(str(doc_path))
    source = source_path.read_text(encoding="utf-8")

    title = find(doc, "Листинг 2 — фрагмент app/Http/Controllers/TicketController.php")
    api_anchor = find(doc, "Отдельным направлением разработки серверной части стала реализация API")

    body = doc._body._element
    children = list(body)
    start = children.index(title._p) + 1
    end = children.index(api_anchor._p)
    for element in children[start:end]:
        body.remove(element)

    listing = "\n".join([
        "<?php",
        "",
        "namespace App\\Http\\Controllers;",
        "",
        "use App\\Models\\Ticket;",
        "use Illuminate\\Database\\Eloquent\\Builder;",
        "use Illuminate\\Http\\RedirectResponse;",
        "use Illuminate\\Http\\Request;",
        "use Illuminate\\View\\View;",
        "",
        "class TicketController extends Controller",
        "{",
        method(source, "    public function index(): View"),
        "",
        method(source, "    public function store(Request $request): RedirectResponse"),
        "",
        method(source, "    private function canEdit(Ticket $ticket): bool"),
        "",
        method(source, "    private function canView(Ticket $ticket): bool"),
        "",
        method(source, "    private function visibleTicketsQuery(): Builder"),
        "}",
    ])
    # Keep the listing concise by replacing long attachment details with an explanatory comment.
    listing = re.sub(
        r"\n\s*if \(\$uploadError = \$this->attachmentUploadError\(\$request\)\) \{.*?\n\s*\}\n\n\s*\$maxAttachmentKb = \$this->attachmentMaxKilobytes\(\);",
        "\n        // Перед сохранением проверяется корректность вложения.\n        $maxAttachmentKb = $this->attachmentMaxKilobytes();",
        listing,
        flags=re.S,
    )
    add_code_after(title, listing)
    doc.save(str(doc_path))
    print(doc_path)


if __name__ == "__main__":
    main()
