from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
}


def text_of(element: ET.Element) -> str:
    return "".join(t.text or "" for t in element.findall(".//w:t", NS))


def paragraph_style_name(p) -> str:
    try:
        return p.style.name
    except Exception:
        return ""


def run_marks(run) -> list[str]:
    marks: list[str] = []
    font = run.font
    try:
        if font.highlight_color:
            marks.append(f"highlight={font.highlight_color}")
    except ValueError:
        highlight = run._element.find(".//w:highlight", NS)
        if highlight is not None:
            key = f"{{{NS['w']}}}val"
            marks.append(f"highlight={highlight.get(key)}")
    if font.color and font.color.rgb:
        marks.append(f"color=#{font.color.rgb}")
    if run.bold:
        marks.append("bold")
    return marks


def summarize_docx(path: Path) -> None:
    print(f"FILE: {path}")
    doc = Document(str(path))
    print("\nHEADINGS:")
    for i, p in enumerate(doc.paragraphs, start=1):
        style = paragraph_style_name(p)
        if style.startswith("Heading") or re.match(r"^\s*(ВВЕДЕНИЕ|ЗАКЛЮЧЕНИЕ|СПИСОК|ПРИЛОЖЕНИЕ|\d+(?:\.\d+)*)", p.text, re.I):
            print(f"{i:04d} [{style}] {p.text[:180]}")

    print("\nMARKED PARAGRAPHS:")
    for i, p in enumerate(doc.paragraphs, start=1):
        marked_runs = []
        for r in p.runs:
            marks = run_marks(r)
            if marks:
                marked_runs.append(f"{'/'.join(marks)}:{r.text!r}")
        if marked_runs:
            print(f"{i:04d} [{paragraph_style_name(p)}] {p.text[:220]}")
            for item in marked_runs[:8]:
                print(f"      {item[:220]}")

    print("\nCOMMENTS:")
    with zipfile.ZipFile(path) as zf:
        names = set(zf.namelist())
        comments_name = "word/comments.xml"
        if comments_name not in names:
            print("no comments.xml")
        else:
            root = ET.fromstring(zf.read(comments_name))
            for c in root.findall(".//w:comment", NS):
                cid = c.attrib.get(f"{{{NS['w']}}}id", "")
                author = c.attrib.get(f"{{{NS['w']}}}author", "")
                text = text_of(c).strip()
                print(f"{cid}: {author}: {text}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("usage: inspect_docx_marks.py file.docx")
    summarize_docx(Path(sys.argv[1]))
