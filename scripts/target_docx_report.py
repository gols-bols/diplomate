from __future__ import annotations

import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "pic": "http://schemas.openxmlformats.org/drawingml/2006/picture",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def para_text(p: ET.Element) -> str:
    return "".join(t.text or "" for t in p.findall(".//w:t", NS)).strip()


def run_text(r: ET.Element) -> str:
    return "".join(t.text or "" for t in r.findall(".//w:t", NS))


def run_mark(r: ET.Element) -> str | None:
    rpr = r.find("w:rPr", NS)
    if rpr is None:
        return None
    marks: list[str] = []
    color = rpr.find("w:color", NS)
    if color is not None:
        val = color.get(f"{{{NS['w']}}}val")
        if val and val.upper() not in {"000000", "111111", "AUTO"}:
            marks.append(f"color={val}")
    highlight = rpr.find("w:highlight", NS)
    if highlight is not None:
        val = highlight.get(f"{{{NS['w']}}}val")
        if val and val != "none":
            marks.append(f"highlight={val}")
    return ", ".join(marks) if marks else None


def make_report(path: Path, out: Path) -> None:
    doc = Document(str(path))
    root = ET.fromstring(zipfile.ZipFile(path).read("word/document.xml"))
    paras = root.findall(".//w:body/w:p", NS)

    lines: list[str] = []
    lines.append(f"FILE: {path}")
    lines.append(f"PARAGRAPHS: {len(doc.paragraphs)}")
    lines.append(f"INLINE_SHAPES: {len(doc.inline_shapes)}")
    lines.append("")
    lines.append("MARKED RUNS:")
    for idx, p in enumerate(paras, start=1):
        ptxt = para_text(p)
        marked = []
        for r in p.findall("w:r", NS):
            mark = run_mark(r)
            txt = run_text(r)
            if mark:
                marked.append(f"{mark}: {txt!r}")
        if marked:
            lines.append(f"{idx:04d}: {ptxt[:240]}")
            lines.extend(f"    {m[:260]}" for m in marked)

    lines.append("")
    lines.append("DRAWINGS:")
    for idx, p in enumerate(paras, start=1):
        if p.findall(".//w:drawing", NS):
            lines.append(f"{idx:04d}: {para_text(p)[:240]}")

    lines.append("")
    lines.append("EMPTY HEADINGS:")
    for i, p in enumerate(doc.paragraphs, start=1):
        style = getattr(p.style, "name", "")
        if style.startswith("Heading") and not p.text.strip():
            lines.append(f"{i:04d}: {style}")

    out.write_text("\n".join(lines), encoding="utf-8")
    print(out)


if __name__ == "__main__":
    make_report(Path(sys.argv[1]), Path(sys.argv[2]))
