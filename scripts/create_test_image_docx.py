"""Один тестовый docx с картинкой — проверка Word/WPS."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

from project_paths import MOCKUP_DIR, DOCS_DIR

out = DOCS_DIR / "ТЕСТ_одна_картинка.docx"
img = MOCKUP_DIR / "mockup_login.png"

doc = Document()
doc.add_paragraph("Если ниже видна картинка — Word отображает PNG из проекта.")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(img), width=Cm(14))
doc.add_paragraph("Рисунок 1").alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.save(str(out))
print("Создан:", out.resolve())
