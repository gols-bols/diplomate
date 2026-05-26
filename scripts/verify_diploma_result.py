from pathlib import Path
from zipfile import ZipFile

from docx import Document

path = Path(__file__).resolve().parent.parent / "docs" / "Диплом_отдельный_финал_60.docx"
doc = Document(path)
with ZipFile(path) as zf:
    print("embedded images:", len([n for n in zf.namelist() if n.startswith("word/media/")]))

print("\n--- 1.4 requirements ---")
for i in range(71, 98):
    p = doc.paragraphs[i]
    t = (p.text or "").strip()
    if not t:
        continue
    pbb = p.paragraph_format.page_break_before
    brs = p._p.xpath(".//w:br")
    page_br = sum(1 for b in brs if b.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type") == "page")
    print(f"{i:3d} pbb={pbb} pageBr={page_br} {t[:90]}")

print("\n--- figures near 2.2 ---")
for i, p in enumerate(doc.paragraphs):
    t = (p.text or "").strip()
    if "2.2.3" in t or t.startswith("Рисунок ") or p._p.xpath(".//w:drawing"):
        print(f"{i:3d} {t[:90] or '[drawing]'}")
