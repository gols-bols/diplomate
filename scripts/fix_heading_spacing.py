"""Remove excessive gaps after headings (empty paragraphs, spacing)."""
from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

from resolve_diploma_v2 import resolve_diploma_v2


def delete_paragraph(p: Paragraph) -> None:
    el = p._element
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def is_heading_paragraph(p: Paragraph) -> bool:
    style = p.style.name if p.style else ""
    text = (p.text or "").strip()
    if style.startswith("Heading"):
        return True
    if text.upper() in {"ВВЕДЕНИЕ", "ЗАКЛЮЧЕНИЕ", "БИБЛИОГРАФИЯ", "СОДЕРЖАНИЕ"}:
        return True
    if re.match(r"^ГЛАВА\s+\d", text, re.I):
        return True
    if re.match(r"^ПРИЛОЖЕНИЕ\s", text, re.I):
        return True
    return False


def format_heading(p: Paragraph) -> None:
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (p.style and p.style.name == "Heading 1") else WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = p.paragraph_format
    fmt.line_spacing = 1.5
    fmt.first_line_indent = Cm(0)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(6)
    fmt.page_break_before = False
    for run in p.runs:
        run.font.name = "Times New Roman"
        try:
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        except Exception:
            pass
        if p.style and p.style.name == "Heading 1":
            run.font.size = Pt(16)
            run.bold = True
        else:
            run.font.size = Pt(14)
            run.bold = True


def format_body_after_heading(p: Paragraph) -> None:
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt = p.paragraph_format
    fmt.line_spacing = 1.5
    fmt.first_line_indent = Cm(1.25)
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.page_break_before = False
    for run in p.runs:
        if run.font.size and run.font.size.pt > 14:
            run.font.size = Pt(14)
        run.font.name = "Times New Roman"


# late import for qn in format_heading
from docx.oxml.ns import qn


def remove_empty_after_headings(doc: Document) -> int:
    removed = 0
    i = 0
    while i < len(doc.paragraphs):
        p = doc.paragraphs[i]
        if is_heading_paragraph(p):
            format_heading(p)
            j = i + 1
            while j < len(doc.paragraphs):
                nxt = doc.paragraphs[j]
                text = (nxt.text or "").strip()
                has_image = bool(nxt._p.xpath(".//w:drawing"))
                if text or has_image:
                    break
                if not text and not has_image:
                    delete_paragraph(nxt)
                    removed += 1
                    continue
                j += 1
            if j < len(doc.paragraphs):
                format_body_after_heading(doc.paragraphs[j])
            i = j
        else:
            i += 1
    return removed


def verify_intro(doc: Document) -> bool:
    for i, p in enumerate(doc.paragraphs):
        if (p.text or "").strip() == "Введение":
            j = i + 1
            while j < len(doc.paragraphs):
                t = (doc.paragraphs[j].text or "").strip()
                if t:
                    return True
                if doc.paragraphs[j]._p.xpath(".//w:drawing"):
                    j += 1
                    continue
                return False
            return False
    return False


def process(path: Path) -> dict:
    doc = Document(str(path))
    removed = remove_empty_after_headings(doc)
    intro_ok = verify_intro(doc)
    doc.save(str(path))
    return {"removed_empty": removed, "intro_ok": intro_ok}


def resolve_final_path() -> Path:
    explicit = Path(r"c:\Users\Admin\Downloads\Диплом_отдельный_расширенный_3_v2_ФИНАЛ.docx")
    if explicit.is_file():
        return explicit.resolve()
    base = resolve_diploma_v2()
    final = base.with_name(base.stem + "_ФИНАЛ.docx")
    if final.is_file():
        return final.resolve()
    return base.resolve()


def main() -> None:
    primary = resolve_final_path()
    base = resolve_diploma_v2()

    backup = primary.with_name(primary.stem + ".bak_spacing.docx")
    shutil.copy2(primary, backup)

    stats = process(primary)

    # sync: base v2 and отдельная копия _СИНХРОН (stem без _ФИНАЛ)
    sync_targets = [base, base.with_name(base.stem + "_СИНХРОН.docx")]
    for target in sync_targets:
        if target.resolve() == primary.resolve():
            continue
        try:
            shutil.copy2(primary, target)
            stats.setdefault("synced", []).append(str(target))
        except PermissionError:
            stats.setdefault("sync_failed", []).append(str(target))

    print("OK", stats)


if __name__ == "__main__":
    main()
