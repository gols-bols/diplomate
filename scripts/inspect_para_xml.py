from pathlib import Path

from docx import Document

doc = Document(Path(__file__).resolve().parent.parent / "docs" / "Диплом_отдельный_финал_60.docx")
for i in range(84, 96):
    el = doc.paragraphs[i]._p
    xml = el.xml if isinstance(el.xml, str) else el.xml.decode("utf-8")
    t = (doc.paragraphs[i].text or "")[:70]
    flags = []
    if "sectPr" in xml:
        flags.append("sectPr")
    if 'w:type="page"' in xml:
        flags.append("pageBr")
    if "pageBreakBefore" in xml:
        flags.append("pbb")
    if "lastRenderedPageBreak" in xml:
        flags.append("lrpb")
    print(i, "|".join(flags) or "-", t)
