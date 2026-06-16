from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.dml import MSO_THEME_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor, Pt
from docx.text.paragraph import Paragraph


BIBLIOGRAPHY = [
    "Об информации, информационных технологиях и о защите информации : Федеральный закон от 27.07.2006 № 149-ФЗ // КонсультантПлюс : [сайт]. – URL: https://www.consultant.ru/document/cons_doc_LAW_61798/ (дата обращения: 08.06.2026).",
    "О персональных данных : Федеральный закон от 27.07.2006 № 152-ФЗ // КонсультантПлюс : [сайт]. – URL: https://www.consultant.ru/document/cons_doc_LAW_61801/ (дата обращения: 08.06.2026).",
    "ГОСТ 19.701-90. Единая система программной документации. Схемы алгоритмов, программ, данных и систем. Условные обозначения и правила выполнения. – Москва : Стандартинформ, 2010. – 26 с.",
    "ГОСТ Р 7.0.100-2018. Система стандартов по информации, библиотечному и издательскому делу. Библиографическая запись. Библиографическое описание. Общие требования и правила составления. – Москва : Стандартинформ, 2018. – 124 с.",
    "ГОСТ Р ИСО/МЭК 25010-2015. Информационные технологии. Системная и программная инженерия. Требования и оценка качества систем и программного обеспечения. Модели качества систем и программных продуктов. – Москва : Стандартинформ, 2015. – 49 с.",
    "Ларман, К. Применение UML 2.0 и шаблонов проектирования : практическое руководство / К. Ларман. – 3-е изд. – Москва : Вильямс, 2019. – 736 с.",
    "Мартин, Р. Чистая архитектура. Искусство разработки программного обеспечения / Р. Мартин. – Санкт-Петербург : Питер, 2018. – 352 с.",
    "Маклафлин, Б. PHP и MySQL. Исчерпывающее руководство / Б. Маклафлин. – Санкт-Петербург : Питер, 2021. – 544 с.",
    "Composer Documentation // Composer : [сайт]. – URL: https://getcomposer.org/doc/ (дата обращения: 08.06.2026).",
    "Git Documentation // Git : [сайт]. – URL: https://git-scm.com/doc (дата обращения: 08.06.2026).",
    "Laravel Documentation // Laravel : [сайт]. – URL: https://laravel.com/docs (дата обращения: 08.06.2026).",
    "MySQL 8.4 Reference Manual // Oracle : [сайт]. – URL: https://dev.mysql.com/doc/refman/8.4/en/ (дата обращения: 08.06.2026).",
    "PHP Manual // PHP : [сайт]. – URL: https://www.php.net/manual/ru/ (дата обращения: 08.06.2026).",
    "Postman Learning Center // Postman : [сайт]. – URL: https://learning.postman.com/docs/ (дата обращения: 08.06.2026).",
]


def iter_all_paragraphs(doc: Document):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def paragraph_after(paragraph: Paragraph, text: str = "", style=None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if style is not None:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def delete_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def find_para(doc: Document, text: str) -> Paragraph:
    for p in doc.paragraphs:
        if text in p.text:
            return p
    raise ValueError(text)


def replace_bibliography(doc: Document) -> None:
    heading = find_para(doc, "Библиография")
    appendix = find_para(doc, "ПРИЛОЖЕНИЕ 1. Код для функционирования системы заявок")
    body = doc._body._element
    children = list(body)
    start = children.index(heading._p) + 1
    end = children.index(appendix._p)
    for element in children[start:end]:
        body.remove(element)

    current = heading
    for number, item in enumerate(BIBLIOGRAPHY, start=1):
        p = paragraph_after(current, f"{number}. {item}")
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.space_after = Pt(0)
        current = p


def force_black_run(run) -> None:
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.color.theme_color = None
    rpr = run._element.get_or_add_rPr()
    color = rpr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        rpr.append(color)
    color.set(qn("w:val"), "000000")
    if color.get(qn("w:themeColor")):
        color.attrib.pop(qn("w:themeColor"), None)


def force_all_text_black(doc: Document) -> None:
    for style in doc.styles:
        if hasattr(style, "font"):
            try:
                style.font.color.rgb = RGBColor(0, 0, 0)
            except Exception:
                pass

    for paragraph in iter_all_paragraphs(doc):
        for run in paragraph.runs:
            force_black_run(run)

    for section in doc.sections:
        for part in (section.header, section.footer):
            for paragraph in part.paragraphs:
                for run in paragraph.runs:
                    force_black_run(run)
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                force_black_run(run)


def main() -> None:
    path = Path(sys.argv[1])
    doc = Document(str(path))
    replace_bibliography(doc)
    force_all_text_black(doc)
    doc.save(str(path))
    print(path)


if __name__ == "__main__":
    main()
