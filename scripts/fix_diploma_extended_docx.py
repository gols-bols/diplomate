#!/usr/bin/env python3
"""Fixes for Диплом_отдельный_расширенный.docx: structure, numbering, formatting."""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

SRC = Path("/Users/gnome/Downloads/Диплом_отдельный_расширенный.docx")
BAK = SRC.with_suffix(".docx.bak")


def scrub_run_xml(run) -> None:
    r = run._element
    rPr = r.find(qn("w:rPr"))
    if rPr is None:
        return
    for tag in ("w:highlight", "w:shd"):
        for el in list(rPr.findall(qn(tag))):
            rPr.remove(el)
    for el in list(rPr.findall(qn("w:color"))):
        rPr.remove(el)


def set_run_black_14(run, bold: bool | None = None) -> None:
    scrub_run_xml(run)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.size = Pt(14)
    if bold is not None:
        run.bold = bold


def iter_all_runs(doc: Document):
    for p in doc.paragraphs:
        for r in p.runs:
            yield r
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        yield r


def move_paragraph_range_after(doc: Document, start: int, end: int, after: int) -> None:
    if start > end or start < 0:
        return
    ps = [doc.paragraphs[i]._p for i in range(start, end + 1)]
    anchor = doc.paragraphs[after]._p
    for p_el in ps:
        parent = p_el.getparent()
        if parent is not None:
            parent.remove(p_el)
    cur = anchor
    for p_el in ps:
        cur.addnext(p_el)
        cur = p_el


def bump_ch2_heading_text(text: str, delta: int = 2) -> str:
    m = re.match(r"^(2\.)((?:\d+\.)*\d+)\s+(.*)$", text.strip())
    if not m:
        return text
    segs = m.group(2).split(".")
    segs[0] = str(int(segs[0]) + delta)
    return m.group(1) + ".".join(segs) + " " + m.group(3)


def shrink_ch1_heading_text(text: str, delta: int = 2) -> str:
    m = re.match(r"^(1\.)((?:\d+\.)*\d+)\s+(.*)$", text.strip())
    if not m:
        return text
    segs = m.group(2).split(".")
    n = int(segs[0])
    if n < 6:
        return text
    segs[0] = str(n - delta)
    return m.group(1) + ".".join(segs) + " " + m.group(3)


def find_h1_index(doc: Document, predicate) -> int:
    for i, p in enumerate(doc.paragraphs):
        if (p.style and p.style.name == "Heading 1") and predicate(p.text.strip()):
            return i
    return -1


def main() -> None:
    shutil.copy2(SRC, BAK)
    doc = Document(str(SRC))

    # Remove editor note from practical significance paragraph
    p39 = doc.paragraphs[39]
    if "Почему анализ предметной области" in p39.text:
        for r in list(p39.runs):
            if "Почему анализ" in r.text:
                r.text = ""

    # Rephrase highlighted fragment in introduction
    p25 = doc.paragraphs[25]
    if "эпизодическом характере неисправностей" in p25.text:
        new = p25.text.replace(
            "при небольшом количестве пользователей и эпизодическом характере неисправностей",
            "при ограниченной численности пользователей и редких инцидентах отказов техники",
        )
        p25.clear()
        r0 = p25.add_run(new)
        r0.font.size = Pt(14)

    # Move old 1.4–1.5 (stack + architecture) to start of chapter 2 (project part)
    start_move, end_move, after_move = 82, 106, 202
    move_paragraph_range_after(doc, start_move, end_move, after_move)

    i_ch1 = find_h1_index(doc, lambda t: t.startswith("1 Аналит") or t.startswith("1 Анализ"))
    i_ch2 = find_h1_index(doc, lambda t: "Практическая" in t or "Проектная" in t)
    i_ch3 = find_h1_index(doc, lambda t: t.startswith("3 "))

    center_titles_exact = {
        "Перечень сокращений и обозначений",
        "Содержание",
        "Введение",
        "Заключение",
        "Библиография",
    }

    # --- Heading 1 ---
    for p in doc.paragraphs:
        if not p.style or p.style.name != "Heading 1":
            continue
        t = p.text.strip()
        if t == "1 Аналитическая часть":
            p.text = "1 Анализ предметной области"
            t = p.text.strip()
        if t.startswith("2.") and "Практическая" in t:
            p.text = "2 Проектная часть"
            t = p.text.strip()
        if t.startswith("2 ") and "Практическая" in t:
            p.text = "2 Проектная часть"
            t = p.text.strip()

        if t in center_titles_exact:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.page_break_before = t in ("Заключение", "Библиография")
            for r in p.runs:
                set_run_black_14(r, bold=True)
                r.font.size = Pt(16)
        elif t.startswith(("1 Анализ", "2 Проектная", "3 ")):
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.page_break_before = True
            for r in p.runs:
                set_run_black_14(r, bold=True)
                r.font.size = Pt(16)
        elif t.startswith("Приложение"):
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.page_break_before = True
            for r in p.runs:
                set_run_black_14(r, bold=True)
                r.font.size = Pt(16)

    # Refresh indices after H1 text changes
    i_ch1 = find_h1_index(doc, lambda t: t.startswith("1 Анализ"))
    i_ch2 = find_h1_index(doc, lambda t: t.startswith("2 Проектная"))
    i_ch3 = find_h1_index(doc, lambda t: t.startswith("3 "))

    # --- Heading 2: numbering + typography ---
    for i, p in enumerate(doc.paragraphs):
        if not p.style or p.style.name != "Heading 2":
            continue
        t = p.text.strip()
        if not t:
            continue

        if i_ch1 != -1 and i_ch2 != -1 and i_ch1 < i < i_ch2:
            if t.startswith("1.4 Выбор технологического стека"):
                continue
            if t.startswith("1.5 Обоснование архитектурных решений"):
                continue
            p.text = shrink_ch1_heading_text(p.text, 2)
            t = p.text.strip()
            if "Выводы по аналитической части" in t:
                p.text = t.replace(
                    "Выводы по аналитической части",
                    "Выводы по результатам анализа предметной области",
                )
            if "Дополнительные выводы по аналитической части" in t:
                p.text = t.replace(
                    "Дополнительные выводы по аналитической части",
                    "Дополнительные выводы по анализу предметной области",
                )

        if i_ch2 != -1 and i_ch3 != -1 and i_ch2 < i < i_ch3:
            if t.startswith("1.4 Выбор технологического стека"):
                p.text = t.replace("1.4 ", "2.1 ", 1)
            elif t.startswith("1.5 Обоснование архитектурных решений"):
                p.text = t.replace("1.5 ", "2.2 ", 1)
            elif t.startswith("2.1 ") and "Выбор технологического стека" in t:
                pass
            elif t.startswith("2.2 ") and "Обоснование архитектурных решений" in t:
                pass
            elif re.match(r"^2\.\d+", t):
                p.text = bump_ch2_heading_text(p.text, 2)

        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.space_before = Pt(6)
        for r in p.runs:
            set_run_black_14(r, bold=True)
            r.font.size = Pt(14)

    for p in doc.paragraphs:
        if p.style and p.style.name == "Heading 3":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.space_before = Pt(6)
            for r in p.runs:
                set_run_black_14(r, bold=True)
                r.font.size = Pt(14)

    # --- Normal paragraphs ---
    for p in doc.paragraphs:
        if not p.style or p.style.name != "Normal":
            continue
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        txt = p.text
        if "аналитической части" in txt:
            txt = txt.replace("аналитической части", "анализа предметной области")
            txt = txt.replace("практической главы диплома", "проектной части дипломной работы")
            if txt != p.text:
                p.clear()
                p.add_run(txt)
        for r in p.runs:
            set_run_black_14(r, bold=False)

    for r in iter_all_runs(doc):
        scrub_run_xml(r)
        try:
            r.font.color.rgb = RGBColor(0, 0, 0)
        except Exception:
            pass

    doc.save(str(SRC))
    print("Saved:", SRC)
    print("Backup:", BAK)


if __name__ == "__main__":
    main()
